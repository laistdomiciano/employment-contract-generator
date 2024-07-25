from flask import Flask, request, render_template, send_file
from docx import Document
import io

app = Flask(__name__)


@app.route('/')
def index():
    return render_template('index.html')


@app.route('/generate', methods=['POST'])
def generate():
    party_a = request.form['party_a']
    party_b = request.form['party_b']
    date = request.form['date']

    # Load template and replace placeholders
    doc = Document('templates/contract_template.docx')
    for paragraph in doc.paragraphs:
        if '[PARTY_A]' in paragraph.text:
            paragraph.text = paragraph.text.replace('[PARTY_A]', party_a)
        if '[PARTY_B]' in paragraph.text:
            paragraph.text = paragraph.text.replace('[PARTY_B]', party_b)
        if '[DATE]' in paragraph.text:
            paragraph.text = paragraph.text.replace('[DATE]', date)

    # Save to a BytesIO object
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    return send_file(buffer, as_attachment=True, download_name='contract.docx',
                     mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document')


if __name__ == '__main__':
    app.run(debug=True)
